"""Labelled evaluation suite (the "Evaluation Dashboard").

Implements the evaluation plan:

* Agent 1 (categorisation & domain) - feedback type and EGSQUAL service
  dimensions vs. validated reference labels: Precision, Recall, Macro F1.
* Agent 2 (severity & impact) - sentiment, severity and priority vs.
  validated reference labels: Precision, Recall, Macro F1. Star ratings are
  used only as a supporting indicator for sentiment.
* Agent 3 (ticket escalation) - ticket-generation and escalation decisions
  vs. reference decisions: Precision, Recall, F1, plus the JSON validity rate
  of the generated tickets.
* Overall framework - Multi-Agent vs. Single LLM vs. Sequential LLM Pipeline
  on the same test set: Macro F1, Escalation F1, processing time and
  estimated cost.

Nothing here invents a score: every metric is computed from human-verified
reference labels the user uploads. Without labels, nothing is reported.
Costs are ESTIMATES (tokens counted from the API responses x configured prices).
"""
from __future__ import annotations

import io
import json
import time
from dataclasses import asdict, dataclass, field
from datetime import date
from typing import Callable, Optional

import pandas as pd
from pydantic import ValidationError

import tools
from evaluation import _row
from schemas import ActionType, EscalationStatus, ReviewInput, SimulatedTicket, WorkflowState

# --------------------------------------------------------------------------- #
# Label specification
# --------------------------------------------------------------------------- #
FEEDBACK_TYPES = ["Bug", "Suggestion", "Inquiry", "Praise", "General Complaint"]
SENTIMENTS = ["Negative", "Neutral", "Positive"]
SEVERITIES = ["None", "Low", "Medium", "High", "Critical"]
PRIORITIES = ["None", "P1", "P2", "P3", "P4"]
YES_NO = ["yes", "no"]
ACTION_CODES = {
    ActionType.INTERNAL_TECHNICAL_TICKET: "TECH_TICKET",
    ActionType.SERVICE_IMPROVEMENT_REQUEST: "SERVICE_IMPROVEMENT",
    ActionType.REQUEST_ADDITIONAL_INFO: "REQUEST_INFO",
    ActionType.FOLLOW_UP_EXISTING_COMPLAINT: "FOLLOW_UP",
    ActionType.ESCALATION_CANDIDATE: "ESCALATION_CANDIDATE",
    ActionType.NO_ACTION: "NO_ACTION",
}
ACTIONS = list(ACTION_CODES.values())

LABEL_SPEC: dict[str, dict] = {
    "feedback_type": {"agent": "Agent1", "values": FEEDBACK_TYPES, "help": "primary feedback type of the review"},
    "dimensions": {"agent": "Agent1", "values": None, "help": "EGSQUAL dimension keys of the issues, ';'-separated (empty cell = not labelled; NONE = no issue)"},
    "sentiment": {"agent": "Agent2", "values": SENTIMENTS, "help": "overall sentiment of the text"},
    "severity": {"agent": "Agent2", "values": SEVERITIES, "help": "overall operational severity (None = no issue)"},
    "priority": {"agent": "Agent2", "values": PRIORITIES, "help": "overall priority (None = no issue)"},
    "action": {"agent": "Agent3", "values": ACTIONS, "help": "recommended operational action"},
    "ticket": {"agent": "Agent3", "values": YES_NO, "help": "should a ticket be generated? (no only when no action is required)"},
    "escalation": {"agent": "Agent3", "values": YES_NO, "help": "is this a candidate for (simulated) external escalation?"},
}
LABEL_COLUMNS = list(LABEL_SPEC)
FAILED = "FAILED"  # prediction placeholder when a system produced no output for a review

SYSTEMS = {
    "multi_agent": "Multi-Agent framework",
    "single_llm": "Single LLM",
    "single_llm_cot": "Single LLM (few-shot + CoT)",
    "sequential": "Sequential LLM pipeline",
    "multi_agent_no_loop": "Multi-Agent w/o review loop (ablation)",
}
SYSTEM_ROLE = {
    "multi_agent": "proposed framework",
    "single_llm": "baseline: one zero-shot call",
    "single_llm_cot": "strong baseline: worked examples + step-by-step reasoning",
    "sequential": "baseline: three fixed calls, no tools / verification / loop",
    "multi_agent_no_loop": "ablation: same framework, Agent 2 -> Agent 1 review loop disabled",
}
# Highlight encoding: the proposed framework in the brand accent, every comparison system in one neutral grey;
# systems are always named on the axis / card, so identity never depends on colour.
PROPOSED_COLOR, OTHER_COLOR = "#1BAF7A", "#8A9690"
SYSTEM_COLORS = {k: (PROPOSED_COLOR if k == "multi_agent" else OTHER_COLOR) for k in SYSTEMS}


# --------------------------------------------------------------------------- #
# Tolerant label normalisation
# --------------------------------------------------------------------------- #
def _blank(v) -> bool:
    return v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() in ("", "nan", "NaN")


def _key(v) -> str:
    return " ".join(str(v).strip().lower().replace("_", " ").replace("-", " ").split())


_SYN = {
    "feedback_type": {"bug": "Bug", "error": "Bug", "خلل": "Bug", "suggestion": "Suggestion", "اقتراح": "Suggestion",
                      "inquiry": "Inquiry", "question": "Inquiry", "استفسار": "Inquiry", "praise": "Praise", "مدح": "Praise",
                      "general complaint": "General Complaint", "complaint": "General Complaint", "شكوى": "General Complaint"},
    "sentiment": {"negative": "Negative", "neg": "Negative", "strongly negative": "Negative", "سلبي": "Negative",
                  "neutral": "Neutral", "mixed": "Neutral", "neutral / mixed": "Neutral", "محايد": "Neutral",
                  "positive": "Positive", "pos": "Positive", "إيجابي": "Positive", "ايجابي": "Positive"},
    "severity": {"none": "None", "no": "None", "n/a": "None", "low": "Low", "medium": "Medium", "high": "High", "critical": "Critical"},
    "priority": {"none": "None", "no": "None", "n/a": "None", "p1": "P1", "1": "P1", "p2": "P2", "2": "P2",
                 "p3": "P3", "3": "P3", "p4": "P4", "4": "P4"},
    "yesno": {"yes": "yes", "y": "yes", "true": "yes", "1": "yes", "نعم": "yes",
              "no": "no", "n": "no", "false": "no", "0": "no", "لا": "no"},
}


def _action_code(v) -> Optional[str]:
    k = _key(v)
    for at, code in ACTION_CODES.items():
        if k in (_key(code), _key(at.value), _key(at.name)):
            return code
    for word, code in [("escalat", "ESCALATION_CANDIDATE"), ("follow", "FOLLOW_UP"), ("information", "REQUEST_INFO"),
                       ("info", "REQUEST_INFO"), ("improvement", "SERVICE_IMPROVEMENT"), ("technical", "TECH_TICKET"),
                       ("tech", "TECH_TICKET"), ("no action", "NO_ACTION"), ("none", "NO_ACTION")]:
        if word in k:
            return code
    return None


def _dimension_keys(v) -> Optional[frozenset]:
    kb = tools.egsqual_kb()
    by_name = {_key(d["name_en"]): k for k, d in kb.items()} | {_key(d["name_ar"]): k for k, d in kb.items()}
    by_name |= {_key(d["name_en"].split("(")[0]): k for k, d in kb.items()}
    out = set()
    for part in str(v).replace("|", ";").replace(",", ";").split(";"):
        p = part.strip()
        if not p:
            continue
        if _key(p) in ("none", "no issue"):
            continue
        key = p.upper().replace(" ", "_").replace("/", "_")
        if key in kb:
            out.add(key)
        elif _key(p) in by_name:
            out.add(by_name[_key(p)])
        else:
            return None  # unknown dimension -> reject the cell instead of guessing
    return frozenset(out)


def normalise_label(field_name: str, v):
    """Map a raw cell to the canonical label, or None if blank/unrecognised."""
    if _blank(v):
        return None
    if field_name == "dimensions":
        return _dimension_keys(v)
    if field_name == "action":
        return _action_code(v)
    table = _SYN["yesno"] if field_name in ("ticket", "escalation") else _SYN.get(field_name, {})
    return table.get(_key(v))


def normalise_reference(ref: pd.DataFrame) -> tuple[dict[str, dict[str, object]], pd.DataFrame]:
    """Return ({review_id: {field: label}}, coverage table). Unrecognised cells are reported, never coerced."""
    if ref is None or "review_id" not in ref.columns:
        raise ValueError("The labels file needs a 'review_id' column")
    labels: dict[str, dict[str, object]] = {}
    cov = []
    for f in LABEL_COLUMNS:
        if f not in ref.columns:
            cov.append({"field": f, "agent": LABEL_SPEC[f]["agent"], "labelled": 0, "unrecognised": 0, "examples_unrecognised": ""})
            continue
        n_ok, bad = 0, []
        for _, row in ref.iterrows():
            raw = row[f]
            if _blank(raw):
                continue
            lab = normalise_label(f, raw)
            if lab is None:
                bad.append(str(raw))
                continue
            labels.setdefault(str(row["review_id"]).strip(), {})[f] = lab
            n_ok += 1
        cov.append({"field": f, "agent": LABEL_SPEC[f]["agent"], "labelled": n_ok, "unrecognised": len(bad),
                    "examples_unrecognised": ", ".join(sorted(set(bad))[:5])})
    return labels, pd.DataFrame(cov)


def reviews_for_labels(ref: pd.DataFrame, dataset: Optional[pd.DataFrame]) -> tuple[list[ReviewInput], list[str]]:
    """Resolve the text of each labelled review: the labels file's own `text` column wins, else the dataset."""
    from data_loader import row_to_review

    ds = dataset.set_index("review_id") if dataset is not None else None
    out, missing = [], []
    for _, row in ref.drop_duplicates("review_id").iterrows():
        rid = str(row["review_id"]).strip()
        if "text" in ref.columns and not _blank(row.get("text")):
            rating = row.get("rating")
            out.append(ReviewInput(review_id=rid, text=str(row["text"]), app_id=str(row.get("app_id", "UNKNOWN") or "UNKNOWN"),
                                   app_name=str(row.get("app_name", row.get("app_id", "UNKNOWN")) or "UNKNOWN"),
                                   rating=None if _blank(rating) else int(float(rating))))
        elif ds is not None and rid in ds.index:
            r = ds.loc[rid]
            r = r.iloc[0] if isinstance(r, pd.DataFrame) else r
            out.append(row_to_review(pd.concat([pd.Series({"review_id": rid}), r])))
        else:
            missing.append(rid)
    return out, missing


# --------------------------------------------------------------------------- #
# Annotation sheet
# --------------------------------------------------------------------------- #
def annotation_sheet(df: pd.DataFrame, n: int = 100, seed: int = 42) -> pd.DataFrame:
    """Stratified random sample by star rating with BLANK label columns.
    Predictions are deliberately not pre-filled, so annotators are not biased by the system."""
    d = df.dropna(subset=["text"])
    if d.empty:
        return pd.DataFrame(columns=["review_id", "app_id", "rating", "text", *LABEL_COLUMNS])
    groups = d.groupby(d["rating"].fillna(-1))
    per = max(1, n // max(groups.ngroups, 1))
    parts = [g.sample(min(len(g), per), random_state=seed) for _, g in groups]
    s = pd.concat(parts)
    if len(s) < n:
        rest = d.drop(s.index)
        s = pd.concat([s, rest.sample(min(len(rest), n - len(s)), random_state=seed)])
    s = s.sample(frac=1, random_state=seed).head(n)
    out = s[["review_id", "app_id", "rating", "text"]].copy()
    for c in LABEL_COLUMNS:
        out[c] = ""
    return out.reset_index(drop=True)


def annotation_guidelines() -> str:
    kb = tools.egsqual_kb()
    lines = ["# Annotation guidelines (reference labels)", "",
             "Label each review by reading the ORIGINAL Arabic text only. Do not look at the system's output.",
             "Leave a cell empty if you cannot decide; empty cells are excluded from that metric.",
             "Ideally two annotators label independently and disagreements are resolved before upload.", ""]
    for f, spec in LABEL_SPEC.items():
        vals = spec["values"] if spec["values"] else [f"{k} ({d['name_en']})" for k, d in kb.items()] + ["NONE"]
        lines.append(f"## {f}  ·  {spec['agent']}")
        lines.append(spec["help"])
        lines.append("Allowed values: " + ", ".join(vals))
        lines.append("")
    lines += ["## Severity guide (operational impact, NOT anger or star rating)",
              "Critical: service blocked AND (widespread or financial / personal-data risk) · High: service blocked or failing repeatedly ·",
              "Medium: task completed with clear friction · Low: minor / cosmetic / suggestion · None: no issue (praise, unrelated text)", "",
              "## Escalation",
              "yes only if the beneficiary already complained to the provider, the complaint is unresolved, the response deadline passed,",
              "and the issue is High or Critical. Escalation is simulated and always needs human approval."]
    return "\n".join(lines)


# --------------------------------------------------------------------------- #
# Review workbook (human verification of drafted labels)
# --------------------------------------------------------------------------- #
UNREVIEWED = "AI-drafted (not reviewed)"
VERIFIED = "human-verified"
_DROPDOWNS = {"feedback_type": FEEDBACK_TYPES, "sentiment": SENTIMENTS, "severity": SEVERITIES, "priority": PRIORITIES,
              "action": ACTIONS, "ticket": YES_NO, "escalation": YES_NO, "label_source": [UNREVIEWED, VERIFIED]}


def build_review_workbook(labels: pd.DataFrame, rationale: dict[str, str] | None = None) -> bytes:
    """Excel workbook for verifying drafted labels: 'Labels' (editable, dropdowns), 'Guide', and a locked
    'AI_original' copy so the dashboard can report how many labels the human reviewer changed."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    rationale = rationale or {}
    cols = ["review_id", "app_id", "rating", "text", *LABEL_COLUMNS, "label_source", "ai_rationale", "reviewer_notes"]
    df = labels.copy()
    for c in cols:
        if c not in df.columns:
            df[c] = ""
    df["label_source"] = UNREVIEWED
    df["ai_rationale"] = df["review_id"].map(lambda r: rationale.get(str(r), ""))
    df = df[cols].fillna("")

    wb = Workbook()
    ws = wb.active
    ws.title = "Labels"
    ws.append(cols)
    for row in df.itertuples(index=False):
        ws.append([str(v) if v is not None else "" for v in row])
    head = PatternFill("solid", fgColor="0E3B2C")
    edit = PatternFill("solid", fgColor="FFF7E0")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = {"review_id": 10, "app_id": 14, "rating": 7, "text": 60, "dimensions": 26, "action": 22,
              "label_source": 24, "ai_rationale": 55, "reviewer_notes": 30}
    n = len(df) + 1
    for i, c in enumerate(cols, start=1):
        letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letter].width = widths.get(c, 15)
        for r in range(2, n + 1):
            cell = ws.cell(row=r, column=i)
            cell.alignment = Alignment(wrap_text=c in ("text", "ai_rationale", "reviewer_notes"), vertical="top",
                                       horizontal="right" if c in ("text", "ai_rationale") else "left",
                                       readingOrder=2 if c in ("text", "ai_rationale") else 0)
            if c in LABEL_COLUMNS or c == "label_source":
                cell.fill = edit
        if c in _DROPDOWNS:
            dv = DataValidation(type="list", formula1='"' + ",".join(_DROPDOWNS[c]) + '"', allow_blank=True)
            dv.add(f"{letter}2:{letter}{n}")
            ws.add_data_validation(dv)
        if c == "dimensions":
            dv = DataValidation(type="custom", formula1="TRUE", allow_blank=True, showInputMessage=True,
                                promptTitle="EGSQUAL", prompt="Keys separated by ; or NONE: " + ", ".join(tools.egsqual_kb()))
            dv.add(f"{letter}2:{letter}{n}")
            ws.add_data_validation(dv)
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions

    g = wb.create_sheet("Guide")
    g.sheet_view.rightToLeft = True
    g.column_dimensions["A"].width = 120
    lines = ["طريقة المراجعة",
             "1. اقرئي النص العربي (عمود text) وقارنيه بالتصنيفات في الخلايا الصفراء.",
             "2. إذا التصنيف صحيح اتركيه. إذا غلط اختاري القيمة الصحيحة من القائمة المنسدلة.",
             "3. عمود ai_rationale يشرح سبب تصنيف الذكاء الاصطناعي، استخدميه للمراجعة السريعة بس لا تعتمدين عليه.",
             "4. بعد ما تراجعين الصف كامل غيّري label_source إلى human-verified.",
             "5. إذا ما قدرتي تقررين، امسحي الخلية (تُستبعد من الحساب) واكتبي السبب في reviewer_notes.",
             "6. احفظي الملف بصيغة xlsx وارفعيه في Evaluation Dashboard. لا تعدلين ورقة AI_original.", ""]
    lines += [ln for ln in annotation_guidelines().splitlines()]
    for ln in lines:
        g.append([ln])
        g.cell(row=g.max_row, column=1).alignment = Alignment(wrap_text=True, horizontal="right" if not ln.isascii() else "left")
    g["A1"].font = Font(bold=True, size=14)

    o = wb.create_sheet("AI_original")
    o.append(["review_id", *LABEL_COLUMNS])
    for row in df[["review_id", *LABEL_COLUMNS]].itertuples(index=False):
        o.append(list(row))
    o.protection.sheet = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_labels_file(data: bytes, filename: str) -> tuple[pd.DataFrame, Optional[pd.DataFrame]]:
    """Read a labels CSV or review workbook. Returns (labels, AI_original or None)."""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, dtype=str, keep_default_na=False, na_values=[""])
        ref = sheets.get("Labels", next(iter(sheets.values())))
        return ref, sheets.get("AI_original")
    return pd.read_csv(io.BytesIO(data), encoding="utf-8-sig", dtype=str, keep_default_na=False, na_values=[""]), None


def review_status(ref: pd.DataFrame) -> dict:
    """How many rows a human has marked as verified (label_source column)."""
    if "label_source" not in ref.columns:
        return {"tracked": False, "rows": len(ref), "verified": None, "unreviewed": None}
    src = ref["label_source"].fillna("").astype(str).str.strip().str.lower()
    verified = int((src == VERIFIED).sum())
    return {"tracked": True, "rows": len(ref), "verified": verified, "unreviewed": len(ref) - verified}


def label_changes(ref: pd.DataFrame, original: pd.DataFrame) -> pd.DataFrame:
    """Per field: how many drafted labels the human reviewer changed (compared after normalisation)."""
    cur, _ = normalise_reference(ref)
    orig, _ = normalise_reference(original)
    rows = []
    for f in LABEL_COLUMNS:
        ids = [rid for rid in orig if f in orig[rid]]
        changed = sum(1 for rid in ids if cur.get(rid, {}).get(f) != orig[rid][f])
        rows.append({"field": f, "drafted": len(ids), "changed_by_reviewer": changed,
                     "changed_%": round(100 * changed / len(ids), 1) if ids else None})
    return pd.DataFrame(rows)


# --------------------------------------------------------------------------- #
# Predictions
# --------------------------------------------------------------------------- #
@dataclass
class Prediction:
    system: str
    review_id: str
    feedback_type: Optional[str] = None
    dimensions: Optional[frozenset] = None
    sentiment: Optional[str] = None
    severity: Optional[str] = None
    priority: Optional[str] = None
    action: Optional[str] = None
    ticket: Optional[str] = None
    escalation: Optional[str] = None
    ticket_json: Optional[str] = None
    evidence: list = field(default_factory=list)
    seconds: float = 0.0
    input_tokens: int = 0
    output_tokens: int = 0
    llm_calls: int = 0
    used_fallback: bool = False
    error: Optional[str] = None

    @property
    def failed(self) -> bool:
        return self.error is not None and self.feedback_type is None

    def get(self, f: str):
        v = getattr(self, f)
        return FAILED if v is None and self.failed else v

    def as_row(self) -> dict:
        d = asdict(self)
        d["dimensions"] = ";".join(sorted(self.dimensions)) if self.dimensions is not None else None
        d["evidence"] = " || ".join(self.evidence)
        return d


def sentiment_class(label: str) -> str:
    """Map Agent 2's sentiment intensity to the 3-class reference scheme."""
    k = _key(label)
    if "negative" in k:
        return "Negative"
    if "positive" in k:
        return "Positive"
    return "Neutral"


def rating_sentiment(rating) -> Optional[str]:
    """Star-rating proxy (1-2 negative, 3 neutral, 4-5 positive) - a supporting indicator only."""
    if _blank(rating):
        return None
    r = int(float(rating))
    return "Negative" if r <= 2 else ("Neutral" if r == 3 else "Positive")


def predictions_from_state(state: WorkflowState, system: str = "multi_agent") -> Prediction:
    a1, a2, a3 = state.agent1, state.agent2, state.agent3
    p = Prediction(system=system, review_id=state.review.review_id,
                   used_fallback=any(e.action == "LLM_FALLBACK" for e in state.decision_log))
    if a1 is None or a2 is None or a3 is None:
        p.error = "; ".join(state.errors) or "workflow incomplete"
        return p
    p.feedback_type = a1.feedback_type
    p.dimensions = frozenset(i.dimension for i in a1.issues)
    p.sentiment = sentiment_class(a2.sentiment_intensity)
    p.severity = a2.overall_severity.value if a2.overall_severity else "None"
    p.priority = a2.overall_priority.value if a2.overall_priority else "None"
    p.action = ACTION_CODES[a3.action]
    p.ticket = "yes" if a3.ticket is not None else "no"
    p.escalation = "yes" if a3.escalation.status == EscalationStatus.CANDIDATE_AWAITING_HUMAN_APPROVAL else "no"
    p.ticket_json = a3.ticket.model_dump_json() if a3.ticket is not None else None
    p.evidence = [i.evidence for i in a1.issues]
    if state.errors:
        p.error = "; ".join(state.errors)
    return p


def ticket_json_validity(js: Optional[str]) -> tuple[bool, str]:
    """A ticket is valid when its JSON parses, validates against the SimulatedTicket schema,
    has every required field filled, uses allowed severity/priority values and keeps the safety flags."""
    if js is None:
        return False, "no ticket"
    try:
        data = json.loads(js)
    except (TypeError, json.JSONDecodeError) as e:
        return False, f"invalid JSON: {e}"
    try:
        t = SimulatedTicket.model_validate(data)
    except ValidationError as e:
        return False, f"schema: {e.errors()[0]['loc']} {e.errors()[0]['msg']}"
    miss = tools.ticket_fields_validator(t.model_dump())["missing_fields"]
    if miss:
        return False, f"missing {miss}"
    if t.severity.split(" ")[0] not in ("Low", "Medium", "High", "Critical", "Not"):
        return False, f"invalid severity '{t.severity}'"
    if t.priority not in ("P1", "P2", "P3", "P4"):
        return False, f"invalid priority '{t.priority}'"
    if not t.simulated or not t.human_approval_required:
        return False, "safety flags missing (simulated / human approval)"
    return True, "valid"


def run_multi_agent(workflow, reviews: list[ReviewInput], reference_date: Optional[date] = None,
                    progress: Optional[Callable[[int, int], None]] = None, system: str = "multi_agent") -> list[Prediction]:
    out = []
    llm = workflow.llm
    for k, r in enumerate(reviews):
        i0, o0, c0 = llm.usage_snapshot()
        t0 = time.perf_counter()
        try:
            state = workflow.run(r, reference_date=reference_date)
            p = predictions_from_state(state, system)
        except Exception as e:  # noqa: BLE001 - record, never fake
            p = Prediction(system=system, review_id=r.review_id, error=f"{type(e).__name__}: {e}")
        p.seconds = time.perf_counter() - t0
        i1, o1, c1 = llm.usage_snapshot()
        p.input_tokens, p.output_tokens, p.llm_calls = i1 - i0, o1 - o0, c1 - c0
        out.append(p)
        if progress:
            progress(k + 1, len(reviews))
    return out


# --------------------------------------------------------------------------- #
# Metrics
# --------------------------------------------------------------------------- #
def _pairs(labels: dict, preds: list[Prediction], f: str):
    t, p = [], []
    for pr in preds:
        ref = labels.get(pr.review_id, {})
        if f in ref:
            t.append(ref[f])
            p.append(pr.get(f))
    return t, p


def multiclass(y_true: list, y_pred: list) -> dict:
    """Per-class P/R/F1 + macro averages + accuracy. FAILED predictions count as wrong but are not a class."""
    classes = sorted((set(y_true) | set(y_pred)) - {FAILED, None})
    rows = []
    for lab in classes:
        tp = sum(t == lab and p == lab for t, p in zip(y_true, y_pred))
        fp = sum(t != lab and p == lab for t, p in zip(y_true, y_pred))
        fn = sum(t == lab and p != lab for t, p in zip(y_true, y_pred))
        rows.append(_row(lab, tp, fp, fn, support=sum(t == lab for t in y_true)))
    table = pd.DataFrame(rows, columns=["label", "precision", "recall", "f1", "support"])
    n = len(y_true)
    return {"n": n, "table": table,
            "macro_precision": float(table["precision"].mean()) if n else None,
            "macro_recall": float(table["recall"].mean()) if n else None,
            "macro_f1": float(table["f1"].mean()) if n else None,
            "accuracy": sum(t == p for t, p in zip(y_true, y_pred)) / n if n else None,
            "confusion": confusion(y_true, y_pred)}


def multilabel(y_true: list, y_pred: list) -> dict:
    """Set-based per-dimension P/R/F1; macro over dimensions and micro over all (review, dimension) pairs."""
    P = [p if isinstance(p, frozenset) else frozenset() for p in y_pred]
    classes = sorted(set().union(*y_true, *P)) if y_true else []
    rows = [_row(lab, sum(lab in t and lab in p for t, p in zip(y_true, P)),
                 sum(lab not in t and lab in p for t, p in zip(y_true, P)),
                 sum(lab in t and lab not in p for t, p in zip(y_true, P)),
                 support=sum(lab in t for t in y_true)) for lab in classes]
    table = pd.DataFrame(rows, columns=["label", "precision", "recall", "f1", "support"])
    micro = _row("MICRO", sum(len(t & p) for t, p in zip(y_true, P)), sum(len(p - t) for t, p in zip(y_true, P)),
                 sum(len(t - p) for t, p in zip(y_true, P)), support=0)
    n = len(y_true)
    return {"n": n, "table": table,
            "macro_precision": float(table["precision"].mean()) if len(table) else None,
            "macro_recall": float(table["recall"].mean()) if len(table) else None,
            "macro_f1": float(table["f1"].mean()) if len(table) else None,
            "micro_f1": micro["f1"] if n else None,
            "exact_match": sum(t == p for t, p in zip(y_true, P)) / n if n else None}


def binary(y_true: list, y_pred: list, positive: str = "yes") -> dict:
    tp = sum(t == positive and p == positive for t, p in zip(y_true, y_pred))
    fp = sum(t != positive and p == positive for t, p in zip(y_true, y_pred))
    fn = sum(t == positive and p != positive for t, p in zip(y_true, y_pred))
    tn = sum(t != positive and p != positive for t, p in zip(y_true, y_pred))
    defined = (tp + fn) > 0  # F1 of the positive class is undefined without reference positives
    r = _row(positive, tp, fp, fn, support=tp + fn)
    return {"n": len(y_true), "tp": tp, "fp": fp, "fn": fn, "tn": tn, "defined": defined,
            "precision": r["precision"] if (tp + fp) else None, "recall": r["recall"] if defined else None,
            "f1": r["f1"] if defined else None,
            "accuracy": (tp + tn) / len(y_true) if y_true else None}


def confusion(y_true: list, y_pred: list) -> pd.DataFrame:
    return (pd.DataFrame({"reference": y_true, "predicted": [p if p is not None else FAILED for p in y_pred]})
            .value_counts().rename("count").reset_index())


def agent1_metrics(labels, preds) -> dict:
    return {"feedback_type": multiclass(*_pairs(labels, preds, "feedback_type")),
            "dimensions": multilabel(*_pairs(labels, preds, "dimensions"))}


def agent2_metrics(labels, preds, reviews: dict[str, ReviewInput] | None = None) -> dict:
    out = {f: multiclass(*_pairs(labels, preds, f)) for f in ("sentiment", "severity", "priority")}
    # star rating as a supporting indicator for sentiment
    reviews = reviews or {}
    t, rp, sp = [], [], []
    for pr in preds:
        rs = rating_sentiment(reviews[pr.review_id].rating) if pr.review_id in reviews else None
        if rs is None:
            continue
        ref = labels.get(pr.review_id, {}).get("sentiment")
        if ref is not None:
            t.append(ref)
            rp.append(rs)
        if pr.sentiment is not None:
            sp.append((pr.sentiment, rs))
    out["rating_baseline"] = multiclass(t, rp) if t else None
    out["rating_agreement"] = (sum(a == b for a, b in sp) / len(sp)) if sp else None
    out["rating_agreement_n"] = len(sp)
    return out


def agent3_metrics(labels, preds) -> dict:
    tickets = [p for p in preds if p.ticket == "yes"]
    checks = [(p.review_id, *ticket_json_validity(p.ticket_json)) for p in tickets]
    return {"ticket": binary(*_pairs(labels, preds, "ticket")),
            "escalation": binary(*_pairs(labels, preds, "escalation")),
            "action": multiclass(*_pairs(labels, preds, "action")),
            "json_valid": sum(1 for c in checks if c[1]), "json_total": len(checks),
            "json_validity_rate": (sum(1 for c in checks if c[1]) / len(checks)) if checks else None,
            "json_checks": pd.DataFrame(checks, columns=["review_id", "valid", "detail"])}


def cost_usd(input_tokens: int, output_tokens: int, prices: tuple) -> Optional[float]:
    pin, pout = prices
    if pin is None or pout is None:
        return None
    return input_tokens / 1e6 * pin + output_tokens / 1e6 * pout


MACRO_FIELDS = ["feedback_type", "dimensions", "sentiment", "severity", "priority"]


def overall_table(labels, results: dict[str, list[Prediction]], prices: dict[str, tuple],
                  reviews: dict[str, ReviewInput] | None = None) -> pd.DataFrame:
    """One row per system: macro F1 per field, their mean, escalation F1, time and estimated cost."""
    rows = []
    for sys_key, preds in results.items():
        a1, a2, a3 = agent1_metrics(labels, preds), agent2_metrics(labels, preds), agent3_metrics(labels, preds)
        per = {"feedback_type": a1["feedback_type"]["macro_f1"], "dimensions": a1["dimensions"]["macro_f1"],
               "sentiment": a2["sentiment"]["macro_f1"], "severity": a2["severity"]["macro_f1"], "priority": a2["priority"]["macro_f1"]}
        avail = [v for v in per.values() if v is not None]
        tin, tout = sum(p.input_tokens for p in preds), sum(p.output_tokens for p in preds)
        cost = cost_usd(tin, tout, prices.get(sys_key, (None, None)))
        n = max(len(preds), 1)
        rows.append({
            "system_key": sys_key, "system": SYSTEMS.get(sys_key, sys_key), "reviews": len(preds),
            **{f"macro_f1_{k}": v for k, v in per.items()},
            "macro_f1_mean": sum(avail) / len(avail) if avail else None,
            "escalation_f1": a3["escalation"]["f1"], "ticket_f1": a3["ticket"]["f1"],
            "json_validity_rate": a3["json_validity_rate"],
            "severe_recall": severe_recall(labels, preds)["recall"],
            "evidence_grounding_rate": grounding(preds, reviews or {})["rate"],
            "reviews_with_unsupported_quote": grounding(preds, reviews or {})["reviews_with_unsupported"],
            "time_total_s": sum(p.seconds for p in preds), "time_per_review_s": sum(p.seconds for p in preds) / n,
            "input_tokens": tin, "output_tokens": tout, "llm_calls": sum(p.llm_calls for p in preds),
            "est_cost_usd": cost, "est_cost_per_review_usd": cost / n if cost is not None else None,
            "failed_reviews": sum(1 for p in preds if p.failed),
            "reviews_with_llm_fallback": sum(1 for p in preds if p.used_fallback),
        })
    return pd.DataFrame(rows)


# --------------------------------------------------------------------------- #
# Evidence that the framework is better: reliability metrics, significance, subgroups
# --------------------------------------------------------------------------- #
SEVERE = {"High", "Critical"}


def severe_recall(labels, preds) -> dict:
    """Of the reviews a human labelled High/Critical, how many did the system also rate High/Critical?"""
    t, p = _pairs(labels, preds, "severity")
    yt = ["yes" if x in SEVERE else "no" for x in t]
    yp = ["yes" if x in SEVERE else "no" for x in p]
    return binary(yt, yp)


def grounding(preds, reviews: dict[str, ReviewInput]) -> dict:
    """Share of evidence quotes that are really in the review (normalised substring, tools.evidence_verifier)."""
    total = ok = 0
    bad_reviews = 0
    for p in preds:
        r = reviews.get(p.review_id)
        if r is None or not p.evidence:
            continue
        res = [tools.evidence_verifier(r.text, q)["verified"] for q in p.evidence if str(q).strip()]
        total += len(res)
        ok += sum(res)
        bad_reviews += any(not x for x in res)
    return {"quotes": total, "verified": ok, "rate": ok / total if total else None, "reviews_with_unsupported": bad_reviews}


def correctness(labels, preds, f: str) -> dict[str, bool]:
    """Per review: is the prediction for field f exactly right? (dimensions: exact set match)."""
    out = {}
    for p in preds:
        ref = labels.get(p.review_id, {})
        if f in ref:
            v = p.get(f)
            if f == "dimensions":
                v = v if isinstance(v, frozenset) else None
            out[p.review_id] = v == ref[f]
    return out


def mcnemar(a: dict[str, bool], b: dict[str, bool]) -> dict:
    """Exact (binomial) McNemar test on paired per-review correctness. b01 = only A right, b10 = only B right."""
    from math import comb

    ids = sorted(set(a) & set(b))
    only_a = sum(1 for i in ids if a[i] and not b[i])
    only_b = sum(1 for i in ids if b[i] and not a[i])
    n = only_a + only_b
    if n == 0:
        pval = 1.0
    else:
        k = min(only_a, only_b)
        pval = min(1.0, 2 * sum(comb(n, j) for j in range(k + 1)) / 2 ** n)
    return {"n": len(ids), "only_a_correct": only_a, "only_b_correct": only_b, "p_value": pval}


def _macro_f1(labels, preds, f: str):
    t, p = _pairs(labels, preds, f)
    if not t:
        return None
    return multilabel(t, p)["macro_f1"] if f == "dimensions" else multiclass(t, p)["macro_f1"]


def _fast_macro_f1(t: list, p: list, multilabel_field: bool) -> Optional[float]:
    """Same macro F1 as multiclass()/multilabel(), without pandas (used inside the bootstrap loop)."""
    if not t:
        return None
    if multilabel_field:
        P = [x if isinstance(x, frozenset) else frozenset() for x in p]
        classes = set().union(*t, *P)
        hit = lambda lab, x: lab in x  # noqa: E731
    else:
        P = p
        classes = (set(t) | set(p)) - {FAILED, None}
        hit = lambda lab, x: x == lab  # noqa: E731
    if not classes:
        return None
    f1s = []
    for lab in classes:
        tp = sum(hit(lab, a) and hit(lab, b) for a, b in zip(t, P))
        fp = sum(not hit(lab, a) and hit(lab, b) for a, b in zip(t, P))
        fn = sum(hit(lab, a) and not hit(lab, b) for a, b in zip(t, P))
        pr = tp / (tp + fp) if tp + fp else 0.0
        rc = tp / (tp + fn) if tp + fn else 0.0
        f1s.append(2 * pr * rc / (pr + rc) if pr + rc else 0.0)
    return sum(f1s) / len(f1s)


def bootstrap_ci(labels, preds, f: str, n_boot: int = 1000, seed: int = 42) -> tuple[Optional[float], Optional[float]]:
    """95% percentile bootstrap CI of macro F1, resampling labelled reviews with replacement."""
    import random

    t, p = _pairs(labels, preds, f)
    if len(t) < 5:
        return None, None
    rng = random.Random(seed)
    idx = range(len(t))
    vals = []
    for _ in range(n_boot):
        sample = [rng.choice(idx) for _ in idx]
        v = _fast_macro_f1([t[i] for i in sample], [p[i] for i in sample], f == "dimensions")
        if v is not None:
            vals.append(v)
    if not vals:
        return None, None
    vals.sort()
    return vals[int(0.025 * len(vals))], vals[min(len(vals) - 1, int(0.975 * len(vals)))]


SIG_FIELDS = ["feedback_type", "dimensions", "sentiment", "severity", "priority", "action", "ticket", "escalation"]


def significance_table(labels, results: dict[str, list[Prediction]], proposed: str = "multi_agent") -> pd.DataFrame:
    """Proposed framework vs every other system, per field: accuracy of both + exact McNemar p-value."""
    if proposed not in results:
        return pd.DataFrame()
    rows = []
    base = results[proposed]
    for sys_key, preds in results.items():
        if sys_key == proposed:
            continue
        for f in SIG_FIELDS:
            a, b = correctness(labels, base, f), correctness(labels, preds, f)
            if not a:
                continue
            m = mcnemar(a, b)
            ids = set(a) & set(b)
            acc_a = sum(a[i] for i in ids) / len(ids) if ids else None
            acc_b = sum(b[i] for i in ids) / len(ids) if ids else None
            verdict = ("not significant" if m["p_value"] >= 0.05 else
                       ("framework better" if m["only_a_correct"] > m["only_b_correct"] else "framework worse"))
            rows.append({"compared_with": SYSTEMS.get(sys_key, sys_key), "field": f, "n": m["n"],
                         "framework_accuracy": acc_a, "other_accuracy": acc_b,
                         "only_framework_right": m["only_a_correct"], "only_other_right": m["only_b_correct"],
                         "p_value": round(m["p_value"], 4), "result (alpha 0.05)": verdict})
    return pd.DataFrame(rows)


def ci_table(labels, results: dict[str, list[Prediction]], fields=MACRO_FIELDS, n_boot: int = 1000) -> pd.DataFrame:
    rows = []
    for sys_key, preds in results.items():
        for f in fields:
            v = _macro_f1(labels, preds, f)
            if v is None:
                continue
            lo, hi = bootstrap_ci(labels, preds, f, n_boot=n_boot)
            rows.append({"system_key": sys_key, "system": SYSTEMS.get(sys_key, sys_key), "field": f,
                         "macro_f1": v, "ci_low": lo, "ci_high": hi})
    return pd.DataFrame(rows)


def subgroups(reviews: dict[str, ReviewInput], labels) -> dict[str, set]:
    """Hard cases, each defined from the text or the human labels only (never from a system's output)."""
    from arabic_text import word_count

    g = {"All labelled": set(), "Saudi dialect": set(), "Multi-issue (2+ dimensions)": set(),
         "Short / vague (<= 5 words)": set(), "Rating contradicts content": set()}
    for rid, lab in labels.items():
        r = reviews.get(rid)
        if r is None:
            continue
        g["All labelled"].add(rid)
        if tools.dialect_lexicon_lookup(r.text)["dialect_terms"]:
            g["Saudi dialect"].add(rid)
        if isinstance(lab.get("dimensions"), frozenset) and len(lab["dimensions"]) >= 2:
            g["Multi-issue (2+ dimensions)"].add(rid)
        if word_count(r.text) <= 5:
            g["Short / vague (<= 5 words)"].add(rid)
        if r.rating is not None and ((r.rating >= 4 and lab.get("severity") in SEVERE | {"Medium"})
                                     or (r.rating <= 2 and lab.get("feedback_type") == "Praise")):
            g["Rating contradicts content"].add(rid)
    return g


def subgroup_table(labels, results: dict[str, list[Prediction]], reviews: dict[str, ReviewInput],
                   fields=("feedback_type", "severity", "action")) -> pd.DataFrame:
    """Accuracy per hard-case subgroup and system (exact match per field, averaged over the given fields)."""
    rows = []
    for gname, ids in subgroups(reviews, labels).items():
        if not ids:
            continue
        for sys_key, preds in results.items():
            sub = [p for p in preds if p.review_id in ids]
            row = {"subgroup": gname, "n": len(ids), "system_key": sys_key, "system": SYSTEMS.get(sys_key, sys_key)}
            accs = []
            for f in fields:
                c = correctness(labels, sub, f)
                row[f"acc_{f}"] = sum(c.values()) / len(c) if c else None
                if c:
                    accs.append(row[f"acc_{f}"])
            row["mean_accuracy"] = sum(accs) / len(accs) if accs else None
            rows.append(row)
    return pd.DataFrame(rows)


def predictions_frame(results: dict[str, list[Prediction]]) -> pd.DataFrame:
    return pd.DataFrame([p.as_row() for preds in results.values() for p in preds])


@dataclass
class EvalRun:
    """What the dashboard keeps in the session after a run."""
    results: dict[str, list[Prediction]] = field(default_factory=dict)
    prices: dict[str, tuple] = field(default_factory=dict)
    engine: str = ""
    notes: list[str] = field(default_factory=list)
