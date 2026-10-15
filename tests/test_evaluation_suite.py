"""Tests of the labelled evaluation suite and the comparison baselines (no network)."""
import json

import pandas as pd
import pytest

import baselines as B
import evaluation_suite as E
from llm import LLMClient, LLMError, StubLLM
from schemas import ReviewInput
from workflow import MultiAgentWorkflow

R1 = ReviewInput(review_id="T1", app_id="Absher", app_name="Absher", rating=1,
                 text="التطبيق يطلع خطأ عند تجديد الاقامة وما قدرت اكمل والرسوم انخصمت")
R2 = ReviewInput(review_id="T2", app_id="Tawakkalna", app_name="Tawakkalna", rating=5, text="تطبيق رائع وسهل الاستخدام شكرا")


# ---------------------------------------------------------------- labels
def test_label_normalisation_is_tolerant_but_never_guesses():
    assert E.normalise_label("feedback_type", " bug ") == "Bug"
    assert E.normalise_label("feedback_type", "complaint") == "General Complaint"
    assert E.normalise_label("severity", "HIGH") == "High"
    assert E.normalise_label("priority", "2") == "P2"
    assert E.normalise_label("action", "Internal technical ticket") == "TECH_TICKET"
    assert E.normalise_label("action", "escalation_candidate") == "ESCALATION_CANDIDATE"
    assert E.normalise_label("escalation", "نعم") == "yes"
    assert E.normalise_label("dimensions", "Efficiency; SYSTEM_AVAILABILITY") == frozenset({"EFFICIENCY", "SYSTEM_AVAILABILITY"})
    assert E.normalise_label("dimensions", "NONE") == frozenset()
    assert E.normalise_label("dimensions", "NOT_A_DIMENSION") is None
    assert E.normalise_label("severity", "") is None
    assert E.normalise_label("sentiment", "angry") is None


def test_reference_coverage_reports_unrecognised_cells():
    ref = pd.DataFrame({"review_id": ["T1", "T2"], "severity": ["High", "banana"], "ticket": ["yes", ""]})
    labels, cov = E.normalise_reference(ref)
    assert labels == {"T1": {"severity": "High", "ticket": "yes"}}
    row = cov.set_index("field").loc["severity"]
    assert row["labelled"] == 1 and row["unrecognised"] == 1 and "banana" in row["examples_unrecognised"]
    with pytest.raises(ValueError):
        E.normalise_reference(pd.DataFrame({"x": [1]}))


def test_annotation_sheet_is_stratified_blank_and_reproducible():
    df = pd.DataFrame({"review_id": [f"R{i}" for i in range(50)], "app_id": "A", "rating": [1 + i % 5 for i in range(50)],
                       "text": ["نص"] * 50})
    a, b = E.annotation_sheet(df, n=20), E.annotation_sheet(df, n=20)
    assert len(a) == 20 and a["review_id"].tolist() == b["review_id"].tolist()
    assert a["rating"].value_counts().min() == 4  # 20 / 5 ratings
    assert (a[E.LABEL_COLUMNS] == "").all().all()
    assert "feedback_type" in E.annotation_guidelines()


# ---------------------------------------------------------------- metrics
def test_multiclass_metrics_by_hand():
    m = E.multiclass(["A", "A", "B", "B"], ["A", "B", "B", E.FAILED])
    t = m["table"].set_index("label")
    assert list(t.index) == ["A", "B"]  # FAILED is not a class
    assert t.loc["A", "precision"] == 1.0 and t.loc["A", "recall"] == 0.5
    assert t.loc["B", "precision"] == 0.5 and t.loc["B", "recall"] == 0.5
    assert m["macro_f1"] == pytest.approx((2 / 3 + 0.5) / 2, abs=1e-3)
    assert m["accuracy"] == 0.5


def test_binary_f1_undefined_without_reference_positives():
    m = E.binary(["no", "no"], ["no", "yes"])
    assert m["defined"] is False and m["f1"] is None and m["fp"] == 1
    m = E.binary(["yes", "no", "yes"], ["yes", "yes", "no"])
    assert m["precision"] == 0.5 and m["recall"] == 0.5 and m["f1"] == 0.5


def test_multilabel_macro_and_micro():
    T = [frozenset({"A"}), frozenset({"A", "B"})]
    P = [frozenset({"A"}), frozenset({"A"})]
    m = E.multilabel(T, P)
    t = m["table"].set_index("label")
    assert t.loc["A", "f1"] == 1.0 and t.loc["B", "f1"] == 0.0
    assert m["macro_f1"] == 0.5 and m["micro_f1"] == pytest.approx(0.8)


def test_ticket_json_validity():
    good = json.dumps({"ticket_id": "X", "ticket_type": "t", "title": "T", "application": "A", "affected_service": "s",
                       "issue_description": "d", "egsqual_dimensions": [], "severity": "High", "priority": "P2",
                       "recommended_action": "a", "justification": "j", "evidence": [], "escalation_status": "NOT_ELIGIBLE",
                       "human_approval_required": True, "original_review_id": "R"})
    assert E.ticket_json_validity(good) == (True, "valid")
    assert not E.ticket_json_validity("{not json")[0]
    assert not E.ticket_json_validity(good.replace('"P2"', '"urgent"'))[0]
    assert not E.ticket_json_validity(good.replace('"title": "T"', '"title": " "'))[0]
    assert not E.ticket_json_validity(json.dumps({"title": "only"}))[0]


def test_cost_estimate_from_tokens():
    assert E.cost_usd(1_000_000, 500_000, (4.0, 20.0)) == pytest.approx(14.0)
    assert E.cost_usd(10, 10, (None, None)) is None


def test_rating_is_only_a_supporting_indicator():
    assert [E.rating_sentiment(x) for x in (1, 2, 3, 4, 5, None)] == ["Negative", "Negative", "Neutral", "Positive", "Positive", None]
    assert E.sentiment_class("Strongly negative") == "Negative" and E.sentiment_class("Neutral / mixed") == "Neutral"


# ---------------------------------------------------------------- systems
def test_multi_agent_predictions_and_overall_table_offline():
    wf = MultiAgentWorkflow(LLMClient(provider="offline"))
    preds = E.run_multi_agent(wf, [R1, R2])
    p1, p2 = preds
    assert p1.feedback_type == "Bug" and p1.ticket == "yes" and p1.escalation == "no"
    assert p2.action == "NO_ACTION" and p2.ticket == "no" and p2.severity == "None"
    assert p1.seconds > 0 and p1.input_tokens == 0 and p1.llm_calls == 0
    assert E.ticket_json_validity(p1.ticket_json)[0]
    labels = {"T1": {"feedback_type": "Bug", "severity": "High", "escalation": "no", "ticket": "yes"},
              "T2": {"feedback_type": "Praise", "severity": "None", "escalation": "no", "ticket": "no"}}
    a3 = E.agent3_metrics(labels, preds)
    assert a3["ticket"]["f1"] == 1.0 and a3["escalation"]["f1"] is None and a3["json_validity_rate"] == 1.0
    tab = E.overall_table(labels, {"multi_agent": preds}, {"multi_agent": (None, None)}).iloc[0]
    assert tab["macro_f1_feedback_type"] == 1.0 and tab["est_cost_usd"] is None and tab["failed_reviews"] == 0


def _single(**kw):
    base = dict(feedback_type="Bug", dimensions=["FULFILLMENT"], sentiment="Negative", severity="High", priority="P2",
                action="TECH_TICKET", escalation_candidate=False,
                ticket=B.BaselineTicket(title="Renewal fails", affected_service="Iqama renewal", issue_description="Error",
                                        severity="High", priority="P2", justification="blocked", evidence=["يطلع خطأ"]))
    base.update(kw)
    return B.SingleLLMOutput(**base)


def test_single_llm_baseline_counts_tokens_and_cost():
    stub = StubLLM({"SingleLLMOutput": [_single()]}, usage_per_call=(1000, 200), prices=(4.0, 20.0))
    [p] = B.run_baseline("single_llm", stub, [R1])
    assert p.feedback_type == "Bug" and p.ticket == "yes" and p.action == "TECH_TICKET"
    assert (p.input_tokens, p.output_tokens, p.llm_calls) == (1000, 200, 1)
    assert E.ticket_json_validity(p.ticket_json)[0]
    assert E.cost_usd(p.input_tokens, p.output_tokens, stub.price_per_mtok()) == pytest.approx(0.008)
    assert stub.calls[0][0] == "SingleLLMOutput"


def test_sequential_pipeline_makes_three_calls():
    stub = StubLLM({"StepClassify": [B.StepClassify(feedback_type="Praise", dimensions=[], issues=[], evidence=[])],
                    "StepAssess": [B.StepAssess(sentiment="Positive", severity="None", priority="None")],
                    "StepDecide": [B.StepDecide(action="NO_ACTION", escalation_candidate=False, ticket=None)]},
                   usage_per_call=(500, 100))
    [p] = B.run_baseline("sequential", stub, [R2])
    assert [c[0] for c in stub.calls] == ["StepClassify", "StepAssess", "StepDecide"]
    assert p.llm_calls == 3 and p.input_tokens == 1500 and p.ticket == "no" and p.ticket_json is None


def test_baseline_failures_are_recorded_not_faked():
    stub = StubLLM({"SingleLLMOutput": [LLMError("boom")]})
    [p] = B.run_baseline("single_llm", stub, [R1])
    assert p.failed and p.get("feedback_type") == E.FAILED and "boom" in p.error
    [q] = B.run_baseline("single_llm", LLMClient(provider="offline"), [R1])
    assert q.failed and "offline" in q.error
    m = E.agent1_metrics({"T1": {"feedback_type": "Bug"}}, [p])["feedback_type"]
    assert m["accuracy"] == 0.0 and m["macro_f1"] == 0.0


# ---------------------------------------------------------------- human review of drafted labels
def test_none_label_survives_csv_reading():
    csv = "review_id,severity,priority\nT1,None,None\nT2,High,\n".encode("utf-8-sig")
    ref, orig = E.read_labels_file(csv, "labels.csv")
    labels, _ = E.normalise_reference(ref)
    assert orig is None and labels["T1"] == {"severity": "None", "priority": "None"} and labels["T2"] == {"severity": "High"}


def test_review_workbook_roundtrip_tracks_human_changes():
    import io

    from openpyxl import load_workbook

    draft = pd.DataFrame({"review_id": ["T1", "T2"], "text": ["أ", "ب"], "severity": ["None", "Low"], "ticket": ["no", "yes"]})
    data = E.build_review_workbook(draft, {"T1": "سبب"})
    ref, orig = E.read_labels_file(data, "review.xlsx")
    assert E.review_status(ref) == {"tracked": True, "rows": 2, "verified": 0, "unreviewed": 2}
    assert ref.loc[0, "ai_rationale"] == "سبب" and ref.loc[0, "severity"] == "None"
    assert E.label_changes(ref, orig)["changed_by_reviewer"].sum() == 0

    wb = load_workbook(io.BytesIO(data))
    ws = wb["Labels"]
    hdr = [c.value for c in ws[1]]
    ws.cell(row=3, column=hdr.index("severity") + 1).value = "Medium"
    for r in (2, 3):
        ws.cell(row=r, column=hdr.index("label_source") + 1).value = E.VERIFIED
    buf = io.BytesIO()
    wb.save(buf)
    ref, orig = E.read_labels_file(buf.getvalue(), "review.xlsx")
    assert E.review_status(ref)["unreviewed"] == 0
    ch = E.label_changes(ref, orig).set_index("field")
    assert ch.loc["severity", "changed_by_reviewer"] == 1 and ch.loc["ticket", "changed_by_reviewer"] == 0
    assert "Guide" in wb.sheetnames and wb["AI_original"].protection.sheet


# ---------------------------------------------------------------- comparison evidence
def test_mcnemar_exact_values():
    a = {f"r{i}": True for i in range(10)}
    b = {f"r{i}": i >= 8 for i in range(10)}  # framework alone right on 8, other alone right on 0
    m = E.mcnemar(a, b)
    assert (m["only_a_correct"], m["only_b_correct"]) == (8, 0)
    assert m["p_value"] == pytest.approx(2 / 2 ** 8)  # 0.0078
    assert E.mcnemar(a, a)["p_value"] == 1.0


def test_severe_recall_and_grounding():
    labels = {"T1": {"severity": "High"}, "T2": {"severity": "Low"}}
    preds = [E.Prediction("s", "T1", severity="Medium", evidence=["التطبيق يطلع خطأ", "كلام مخترع غير موجود"]),
             E.Prediction("s", "T2", severity="Low", evidence=[])]
    assert E.severe_recall(labels, preds)["recall"] == 0.0
    g = E.grounding(preds, {"T1": R1, "T2": R2})
    assert g["quotes"] == 2 and g["verified"] == 1 and g["reviews_with_unsupported"] == 1


def test_significance_and_subgroups_and_ci():
    labels = {"T1": {"feedback_type": "Bug", "severity": "High", "dimensions": frozenset({"A", "B"})},
              "T2": {"feedback_type": "Praise", "severity": "None"}}
    res = {"multi_agent": [E.Prediction("multi_agent", "T1", feedback_type="Bug"), E.Prediction("multi_agent", "T2", feedback_type="Praise")],
           "single_llm": [E.Prediction("single_llm", "T1", feedback_type="Bug"), E.Prediction("single_llm", "T2", feedback_type="Bug")]}
    sig = E.significance_table(labels, res).set_index("field")
    assert sig.loc["feedback_type", "only_framework_right"] == 1 and sig.loc["feedback_type", "result (alpha 0.05)"] == "not significant"
    groups = E.subgroups({"T1": R1, "T2": R2}, labels)
    assert groups["Multi-issue (2+ dimensions)"] == {"T1"} and groups["All labelled"] == {"T1", "T2"}
    assert E.bootstrap_ci(labels, res["multi_agent"], "feedback_type") == (None, None)  # too few reviews -> no CI
    sub = E.subgroup_table(labels, res, {"T1": R1, "T2": R2})
    assert set(sub["system_key"]) == {"multi_agent", "single_llm"}


def test_cot_baseline_and_ablation_system():
    out = B.SingleLLMCoTOutput(reasoning="step by step", feedback_type="Praise", sentiment="Positive", severity="None",
                               priority="None", action="NO_ACTION", escalation_candidate=False, evidence=["تطبيق رائع"])
    stub = StubLLM({"SingleLLMCoTOutput": [out]})
    [p] = B.run_baseline("single_llm_cot", stub, [R2])
    assert p.system == "single_llm_cot" and p.action == "NO_ACTION" and p.evidence == ["تطبيق رائع"]
    assert list(B.SingleLLMCoTOutput.model_fields)[0] == "reasoning"  # reasoning is generated before the answer
    preds = E.run_multi_agent(MultiAgentWorkflow(LLMClient(provider="offline"), max_revisions=0), [R1], system="multi_agent_no_loop")
    assert preds[0].system == "multi_agent_no_loop" and preds[0].evidence
